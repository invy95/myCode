# 功能：原简历全部条目（含做市/CTA/研报复现）+ 正知，用紧凑单栏排版，验证能否装进两页。
# 更新：2026-08-30：不删原项目，只收行距、字号和页边距。

"""输出 殷维_简历_原结构_紧排.docx。公司名、职位名、日期、原项目数字不改。"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path(__file__).resolve().parent / "殷维_简历_原结构_紧排.docx"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GRAY = RGBColor(0x55, 0x55, 0x55)


def set_run_font(run, name="微软雅黑", size=10, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color is not None:
        run.font.color.rgb = color


def tight(p, before=0, after=1, line=1.02):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line


def add_center(doc, text, size, bold=False, after=1, color=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tight(p, 0, after, 1.0)
    set_run_font(p.add_run(text), size=size, bold=bold, color=color)


def add_section(doc, text):
    p = doc.add_paragraph()
    tight(p, 5, 2, 1.0)
    set_run_font(p.add_run(text), size=11, bold=True, color=NAVY)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F3A5F")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_line(doc, text, size=10, bold=False, before=0, after=1, color=None):
    p = doc.add_paragraph()
    tight(p, before, after)
    set_run_font(p.add_run(text), size=size, bold=bold, color=color)


def add_job(doc, company, dates, title, bullets):
    p = doc.add_paragraph()
    tight(p, 3, 0)
    set_run_font(p.add_run(company), size=10, bold=True)
    set_run_font(p.add_run("  |  " + title), size=10, color=GRAY)
    set_run_font(p.add_run("    " + dates), size=9.5, color=GRAY)
    for i, b in enumerate(bullets, start=1):
        add_line(doc, f"{i}、{b}", size=10, after=0)


def add_project(doc, name, dates, desc, duties, result):
    p = doc.add_paragraph()
    tight(p, 3, 0)
    set_run_font(p.add_run("● " + name), size=10, bold=True)
    set_run_font(p.add_run("    " + dates), size=9.5, color=GRAY)
    add_line(doc, "项目描述：" + desc, size=10, after=0)
    for i, b in enumerate(duties, start=1):
        add_line(doc, f"{i}、{b}", size=10, after=0)
    add_line(doc, "项目业绩：" + result, size=10, after=1)


def main():
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Cm(1.15)
        sec.bottom_margin = Cm(1.15)
        sec.left_margin = Cm(1.5)
        sec.right_margin = Cm(1.5)

    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(10)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    add_center(doc, "殷  维", 16, bold=True, after=2, color=NAVY)
    add_center(
        doc,
        "女  |  31岁  |  中共党员  |  湖北武汉  |  165cm  |  金融硕士",
        9.5,
        after=0,
        color=GRAY,
    )
    add_center(
        doc,
        "15671678098  |  yw43@foxmail.com  |  武汉市洪山区",
        9.5,
        after=0,
        color=GRAY,
    )
    add_center(doc, "求职意向：AI产品经理（交易创新方向）", 10, bold=True, after=2)

    add_section(doc, "工  作  经  历")
    add_job(
        doc,
        "湖北正知资产管理有限公司（私募，中基协 P1007881）",
        "2026.06-至今",
        "AI投研 / 量化研究",
        [
            "对比 WorkBuddy、Coze 等工具，部署投研智能体并对接微信、飞书，先打通可用流程再迭代。",
            "为看盘同事做盘中放量扫描，并按上午 / 午盘 / 下午 / 收盘推送股票池要点。",
            "把「强势股放量」收成带过滤条件的可回测规则（如收盘价 / 前20日最低价 < 130%），对照回测观察最大单笔亏损变化，并记录低波动、下跌市失效情况。",
            "用 Linux 定时任务保证工作日流程可跑。",
        ],
    )
    add_job(
        doc,
        "量化研究工作室",
        "2024.03-2026.04",
        "二级市场研究员",
        [
            "结合宏观与行业研究框架，研究热点赛道轮动，为策略定调提供参考。",
            "搭建标准化财务选股模型，构建季度更新核心股票池，精准筛选优质个股标的。",
            "结合短线指标，完善个股买卖点参考逻辑。",
            "持续迭代优化投研体系，助力整体策略优化。",
        ],
    )
    add_job(
        doc,
        "江西铜业产融控股有限公司",
        "2023.10-2024.02",
        "投资管理岗",
        [
            "使用 tbquant、聚宽、掘金、文华财经等常见量化平台进行期货、股票因子挖掘、策略开发及回测。",
            "对策略回测结果表现进行分析优化。",
            "根据基金经理要求，使用 python 对券商金工等策略研报进行复现。",
        ],
    )
    add_job(
        doc,
        "量盈私募基金管理有限公司",
        "2022.01-2023.09",
        "量化研究员",
        [
            "进行期货、股票因子挖掘、策略开发及回测。",
            "对策略回测结果表现进行分析优化。",
            "根据基金经理要求，使用 python 对券商金工等策略研报进行复现。",
        ],
    )
    add_job(
        doc,
        "深圳引力波量化科技有限公司",
        "2020.08-2021.12",
        "量化研究员",
        [
            "交易数据获取，熟练使用 Pandas、Numpy、Talib 等 package 对数据进行清洗、分析。",
            "量化策略开发等工作，包括套利、做市等方向，日常优化、维护交易策略。",
            "搭建交易风控体系，严格控制交易风险。",
        ],
    )
    add_job(
        doc,
        "中国平安财产保险股份有限公司",
        "2019.07-2020.07",
        "再保险业务管理岗（管培生）",
        [
            "熟练使用 excel 对再保人资信进行月度更新、季度坏账计提、监控风险指标。",
            "根据再保险资信管理制度，通过市场及同业信息，审核交易对手资信情况。",
            "统计汇总不同业务场景分类下关联交易、非关联交易情况数据汇总。",
        ],
    )

    add_section(doc, "教  育  背  景")
    add_line(doc, "2017.09-2019.06    华中科技大学    金融    硕士（全日制）", size=10, before=2, after=0)
    add_line(doc, "2013.09-2017.06    武汉科技大学    机械工程    本科（全日制）", size=10, after=1)

    add_section(doc, "项  目  经  历")
    add_project(
        doc,
        "交易侧 AI 投研辅助",
        "2026.06-至今",
        "在私募投研场景中，用 AI 工具把盘中看盘信息和公告采集做成可跑通的小流程，先给内部同事用，再用对照回测验证规则，不先上重系统。",
        [
            "对比 WorkBuddy、Coze，部署投研智能体并接到微信、飞书。",
            "做盘中放量扫描和股票池分时段推送（上午 / 午盘 / 下午 / 收盘）。",
            "把「找强势股」写成带过滤条件的规则（收盘价 / 前20日最低价 < 130%），做对照回测。",
            "把「盯公告 / 净利润」拆成采集、校验、失败换源。公开 demo：github.com/invy95/myCode（langgraph-announcement-agent）。",
        ],
        "对照回测中最大单笔亏损约从 -27% 降至 -17%，并写明低波动、下跌市信号变差；公告流程不把模型输出当成成交指令。",
    )
    add_project(
        doc,
        "股票投资研究",
        "2024.03-至今",
        "使用 cursor 等 AI Agent 搭建多维度股票市场分析框架，从基本面、情绪面、技术面等分析市场及个股，包括但不限于实时追踪热点板块、市场整体情绪周期识别、基本面财报数据整理、业绩归因等，辅助投资决策。",
        [
            "搭建财报数据收集系统：使用 tushare 收集全市场个股基本面财报数据并计算各基本面相关因子值。",
            "编写业绩预报、快报爬虫，及时爬取个股相关业绩信息。",
            "设计实时热点概念看板，追踪热点板块。",
        ],
        "对上述股票数据进行收集整理，观测有效性的同时加深了对个股财报基本面、短线技术形态、和市场情绪周期和盘感的理解，开发过程中熟练掌握目前市场上 AI 工具的各个功能，提升工作效率。",
    )
    add_project(
        doc,
        "股票短线多头策略",
        "2022.07-2023.02",
        "使用掘金平台开发股票短线多头策略，通过分钟级别板块因子组合日线级别量价因子进行择时选股获得超额收益，并进行策略实盘，策略开发过程中共挖掘 6 个因子，其中 3 个组合成策略有效性最佳。",
        [
            "通过分析量价关系，组合合适的择时、选股因子，结合入场、出场方式的选择，研发股票策略。",
            "对策略结果进行回测，对参数、因子组合进行优化。",
            "编写实盘策略版本，接入掘金每日更新数据。",
        ],
        "回测样本内最大回撤 2%，年化收益 17.14%，夏普比率 3.4。",
    )
    add_project(
        doc,
        "金融工程研报复现",
        "2022.11-2023.01",
        "使用 python 结合 wind 数据库对金融工程研报《3M 板块轮动策略》进行复现。",
        [
            "对研报进行学习研究，使用 wind 收集研报中涉及的微观部分数据。",
            "对策略微观因子设计算法进行计算，完成因子面板。",
        ],
        "完成微观因子的数据收集和因子时间截面的输出。",
    )
    add_project(
        doc,
        "期货 CTA 策略研究",
        "2022.01-2022.03",
        "使用 python 结合 tbquant 针对“基于期货的模式识别”进行策略开发及优化工作。",
        [
            "使用 TBQuant 平台和 python 对接，利用各自功能优势实现数据处理并传递。",
            "基于拓扑空间结构距离开发模式识别策略算法开发及细化成型。",
            "模式识别策略应用于股指期货市场的回测情况及结果分析。",
            "基于股指期货的模式识别策略优化及参数优化。",
        ],
        "基于股指期货回测 2012-2018 年连续 7 年每年回测收益率为正，年化收益率 80%。",
    )
    add_project(
        doc,
        "做市项目",
        "2020.12-2021.12",
        "为交易所提供流动性，积累了丰富的风控经验。",
        [
            "对接项目方了解需求，并为交易所编写做市策略提供流动性。",
            "设置风控系统，包括：物理风控及策略程序风控等。",
            "盯盘控制极端行情下的交易风险，及时进行手动反向平仓操作。",
        ],
        "提升活跃交易所交投量。",
    )

    add_section(doc, "荣  誉  技  能")
    for line in [
        "证券从业资格证书、基金从业资格证书、银行从业资格证书；计算机二级；英语四、六级（558）",
        "熟练掌握 python、matlab、wind、Stata、Excel、Access、PPT；能用 wind 结合 excel 做金融数据分析",
        "硕士论文：《医药行业上市公司价值评估研究--以云南白药为例》",
        "熟练使用 claude、cursor、Coze 等 AI 工具辅助编程，了解 MCP，能搭建 AI agent team；能把投研任务拆成采集、校验、失败换源，不把模型输出当成下单指令",
        "近期：Xpert 金融 Skill 评测、TalentsAI Agentic Coding × 金融已过审（远程交付，非全职经历）",
    ]:
        add_line(doc, line, size=10, after=0)

    doc.save(OUT)
    print(f"已生成：{OUT}")


if __name__ == "__main__":
    main()
