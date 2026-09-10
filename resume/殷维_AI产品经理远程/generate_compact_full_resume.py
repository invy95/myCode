# 功能：按 BOSS「AI产品经理（远程｜投研与资产配置）」改写原结构紧排简历。
# 更新：2026-09-06：只换意向和投研/配置相关表述；公司名、职位名、日期、项目数字与原 PDF 一致。
# 更新：2026-09-07：按本人要求，量化研究工作室结束月由 2026.04 改为 2026.06。
# 更新：2026-09-07：投递远程投研岗，正知职务名由「AI投研 / 量化研究」改为「AI投研」。

"""输出 殷维_简历_原结构_紧排.docx。目标 Word 下两页。"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path(__file__).resolve().parent / "殷维_简历_原结构_紧排.docx"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GRAY = RGBColor(0x55, 0x55, 0x55)


def set_run_font(run, name="微软雅黑", size=10, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color is not None:
        run.font.color.rgb = color


def tight(p, before=0, after=0, line=1.05):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line


def add_center(doc, text, size, bold=False, after=0, color=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tight(p, 0, after)
    set_run_font(p.add_run(text), size=size, bold=bold, color=color)


def add_section(doc, text):
    p = doc.add_paragraph()
    tight(p, 5, 1)
    set_run_font(p.add_run(text), size=11, bold=True, color=NAVY)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F3A5F")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_line(doc, text, size=10, bold=False, before=0, after=0, color=None):
    p = doc.add_paragraph()
    tight(p, before, after)
    set_run_font(p.add_run(text), size=size, bold=bold, color=color)


def add_entry_head(doc, title, dates, subtitle=None):
    p = doc.add_paragraph()
    tight(p, 4, 0)
    set_run_font(p.add_run(title), size=10, bold=True)
    if subtitle:
        set_run_font(p.add_run("  |  " + subtitle), size=10, color=GRAY)
    set_run_font(p.add_run("    " + dates), size=9.5, color=GRAY)


def add_job(doc, company, dates, title, bullets):
    add_entry_head(doc, company, dates, title)
    for i, b in enumerate(bullets, start=1):
        add_line(doc, f"{i}、{b}")


def add_project(doc, name, dates, desc, duties, result):
    add_entry_head(doc, "● " + name, dates)
    add_line(doc, "项目描述：" + desc)
    for i, b in enumerate(duties, start=1):
        add_line(doc, f"{i}、{b}")
    add_line(doc, "项目业绩：" + result)


def main():
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Cm(1.3)
        sec.bottom_margin = Cm(1.3)
        sec.left_margin = Cm(1.6)
        sec.right_margin = Cm(1.6)

    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(10)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    add_center(doc, "殷  维", 16, bold=True, after=1, color=NAVY)
    add_center(
        doc,
        "女  |  31岁  |  中共党员  |  湖北武汉  |  金融硕士  |  15671678098  |  yw43@foxmail.com",
        9.5,
        color=GRAY,
    )
    add_center(doc, "求职意向：AI产品经理（远程｜投研与资产配置）", 10, bold=True, after=1)

    add_section(doc, "工  作  经  历")
    add_job(
        doc,
        "湖北正知资产管理有限公司（私募，中基协 P1007881）",
        "2026.06-至今",
        "AI投研",
        [
            "对比 WorkBuddy、Coze 等工具，在投研场景部署智能体并对接微信、飞书，先打通可用流程再迭代；Linux 定时任务保证工作日可跑。",
            "给研究同事做市场数据解析：盘中放量扫描，并按上午 / 午盘 / 下午 / 收盘推送股票池要点，辅助盘中判断。",
            "把「强势股放量」收成带过滤的可回测规则（收盘价 / 前20日最低价 < 130%），对照回测看最大单笔亏损变化，并记录低波动、下跌市失效情况。",
        ],
    )
    add_job(
        doc,
        "量化研究工作室",
        "2024.03-2026.06",
        "二级市场研究员",
        [
            "结合宏观与行业研究框架，研究热点赛道轮动，为策略定调提供逻辑梳理与参考。",
            "搭建标准化财务选股模型，构建季度更新核心股票池；结合短线指标完善个股买卖点参考逻辑。",
        ],
    )
    add_job(
        doc,
        "江西铜业产融控股有限公司",
        "2023.10-2024.02",
        "投资管理岗",
        [
            "使用 tbquant、聚宽、掘金、文华财经等量化平台进行期货、股票因子挖掘、策略开发及回测，并对回测表现进行分析优化。",
        ],
    )
    add_job(
        doc,
        "量盈私募基金管理有限公司",
        "2022.01-2023.09",
        "量化研究员",
        [
            "进行期货、股票因子挖掘、策略开发及回测，并对策略回测结果表现进行分析优化。",
            "根据基金经理要求，使用 python 对券商金工等策略研报进行复现。",
        ],
    )
    add_job(
        doc,
        "深圳引力波量化科技有限公司",
        "2020.08-2021.12",
        "量化研究员",
        [
            "直接对接交易所需求方，把「盘口要铺得像真实交易」这类模糊要求，落成可执行参数：买一至买五各档的价位排布、每档挂单量、后台逐档铺单的秒级间隔。",
            "自己用 python 在交易所接口上实现这套铺单程序并上线，搭建物理风控与策略程序风控，盯盘处理极端行情下的撤单与反向平仓。",
            "交易数据获取与清洗（Pandas、Numpy、Talib），支持套利、做市方向的策略开发与维护。",
        ],
    )
    add_job(
        doc,
        "中国平安财产保险股份有限公司",
        "2019.07-2020.07",
        "再保险业务管理岗（管培生）",
        [
            "使用 excel 对再保人资信进行月度更新、季度坏账计提、监控风险指标，并审核交易对手资信、汇总关联与非关联交易数据。",
        ],
    )

    add_section(doc, "教  育  背  景")
    add_line(doc, "2017.09-2019.06    华中科技大学    金融    硕士（全日制）", before=1)
    add_line(doc, "2013.09-2017.06    武汉科技大学    机械工程    本科（全日制）")

    add_section(doc, "项  目  经  历")
    add_project(
        doc,
        "投研侧 AI：意图落地与决策辅助",
        "2026.06-至今",
        "用 AI 把研报、公告和市场扫描做成可跑通的小流程，把同事的投资意图收成带过滤的规则和推送，先内部用、再对照回测验证，不先上重系统。",
        [
            "对比 WorkBuddy、Coze 部署投研智能体并接到微信、飞书；做市场数据解析（盘中放量扫描）与股票池分时段推送。",
            "把「找强势股」写成带风控过滤的规则（收盘价 / 前20日最低价 < 130%）并对照回测；公告采集拆成采集、校验、失败换源。demo：github.com/invy95/myCode（langgraph-announcement-agent）。",
        ],
        "对照回测中最大单笔亏损约从 -27% 降至 -17%，并写明低波动、下跌市信号变差；不把模型输出当成成交或配置指令。",
    )
    add_project(
        doc,
        "股票投资研究",
        "2024.03-至今",
        "使用 cursor 等 AI Agent 搭建多维度股票分析框架，从基本面、情绪面、技术面解析市场数据、梳理逻辑、整理财报与业绩归因，辅助策略研究决策。",
        [
            "搭建财报数据收集系统：用 tushare 收集全市场个股基本面数据并计算相关因子值。",
            "编写业绩预报、快报爬虫及时爬取个股业绩信息；设计实时热点概念看板。",
        ],
        "加深对财报基本面、短线技术形态与市场情绪周期的理解，能把研究想法落到可核对的选股条件。",
    )
    add_project(
        doc,
        "股票短线多头策略",
        "2022.07-2023.02",
        "使用掘金平台开发股票短线多头策略，通过分钟级别板块因子组合日线级别量价因子进行择时选股获得超额收益，并进行策略实盘，共挖掘 6 个因子，其中 3 个组合成策略有效性最佳。",
        [
            "通过分析量价关系组合择时、选股因子，结合入场出场方式研发策略，并对参数、因子组合进行回测优化。",
            "编写实盘策略版本，接入掘金每日更新数据。",
        ],
        "回测样本内最大回撤 2%，年化收益 17.14%，夏普比率 3.4。",
    )
    add_project(
        doc,
        "金融工程研报复现",
        "2022.11-2023.01",
        "使用 python 结合 wind 数据库对金融工程研报《3M 板块轮动策略》进行复现。",
        [
            "学习研究研报并用 wind 收集涉及的微观部分数据，对策略微观因子设计算法进行计算。",
        ],
        "完成微观因子的数据收集和因子时间截面的输出。",
    )
    add_project(
        doc,
        "期货 CTA 策略研究",
        "2022.01-2022.03",
        "使用 python 结合 tbquant 针对“基于期货的模式识别”进行策略开发及优化工作。",
        [
            "使用 TBQuant 平台与 python 对接实现数据处理并传递；基于拓扑空间结构距离开发模式识别策略算法。",
            "将模式识别策略应用于股指期货市场，进行回测结果分析与策略、参数优化。",
        ],
        "基于股指期货回测 2012-2018 年连续 7 年每年回测收益率为正，年化收益率 80%。",
    )
    add_project(
        doc,
        "做市系统：需求对接 → 铺单方案设计 → 上线",
        "2020.12-2021.12",
        "为交易所提供流动性。从需求对接、方案设计到程序实现与上线盯盘，全流程由我一人完成。",
        [
            "与交易所对接明确流动性目标与盘口形态要求，把「看起来要像真实挂单」翻译成可落地参数：五档价位如何排布、每档挂多少、按几秒节奏逐档铺出。",
            "设计后台铺单序列并用程序实现上线；配套物理风控与程序风控，盯盘控制极端行情风险并及时手动反向平仓。",
        ],
        "提升交易所交投活跃度；完整跑通「需求 → 参数设计 → 实现 → 上线盯盘」闭环。",
    )

    add_section(doc, "荣  誉  技  能")
    for line in [
        "证券、基金、银行从业资格证书；计算机二级；英语四、六级（558）。硕士论文：《医药行业上市公司价值评估研究--以云南白药为例》",
        "熟练掌握 python、matlab、wind、Stata、Excel、Access、PPT，能用 wind 结合 excel 做金融数据分析",
        "熟练使用 claude、cursor、Coze，了解 MCP，能搭 agent team；清楚大模型会编、会过时，任务拆成采集、校验、失败换源，模型输出不当成交或配置指令",
        "近期：Xpert 金融 Skill 评测、TalentsAI Agentic Coding × 金融已过审（远程交付，非全职经历）",
    ]:
        add_line(doc, line)

    doc.save(OUT)
    print(f"已生成：{OUT}")


if __name__ == "__main__":
    main()
