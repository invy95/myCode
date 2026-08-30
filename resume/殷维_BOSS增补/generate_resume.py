# 功能：把殷维 BOSS 增补简历渲染成可投递的 docx。
# 更新：2026-08-30：按 11 号 PDF 时间线 + 正知/AI 投研增补生成一页半中文简历。

"""读取结构化段落，输出 resume/殷维_BOSS增补/殷维_简历_BOSS增补.docx。"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path(__file__).resolve().parent / "殷维_简历_BOSS增补.docx"


def set_run_font(run, name="微软雅黑", size=10.5, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color is not None:
        run.font.color.rgb = color


def add_heading_bar(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    set_run_font(run, size=11, bold=True, color=RGBColor(0x1F, 0x3A, 0x5F))
    return p


def add_job(doc, title, dates, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    r1 = p.add_run(title)
    set_run_font(r1, size=10.5, bold=True)
    r2 = p.add_run("    " + dates)
    set_run_font(r2, size=10, color=RGBColor(0x55, 0x55, 0x55))
    for b in bullets:
        bp = doc.add_paragraph(style="List Bullet")
        bp.paragraph_format.space_before = Pt(0)
        bp.paragraph_format.space_after = Pt(1)
        bp.paragraph_format.line_spacing = 1.05
        bp.paragraph_format.left_indent = Cm(0.6)
        run = bp.add_run(b)
        set_run_font(run, size=10)


def main():
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Cm(1.4)
        sec.bottom_margin = Cm(1.4)
        sec.left_margin = Cm(1.6)
        sec.right_margin = Cm(1.6)

    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(10.5)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(2)
    r = name.add_run("殷  维")
    set_run_font(r, size=18, bold=True, color=RGBColor(0x1F, 0x3A, 0x5F))

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(1)
    r = sub.add_run("量化研究  /  AI投研应用  |  华中科技大学 · 金融硕士")
    set_run_font(r, size=10.5, bold=True)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(2)
    r = contact.add_run(
        "女  |  31岁  |  中共党员  |  武汉（可到岗北京/上海）  |  15671678098  |  yw43@foxmail.com"
    )
    set_run_font(r, size=9.5, color=RGBColor(0x44, 0x44, 0x44))

    intent = doc.add_paragraph()
    intent.paragraph_format.space_after = Pt(2)
    r = intent.add_run("求职意向：")
    set_run_font(r, size=10.5, bold=True)
    r = intent.add_run("量化研究 / AI投研应用（数据采集与校验、投研工作流、策略回测）")
    set_run_font(r, size=10.5)

    add_heading_bar(doc, "核心匹配")
    for line in [
        "金融硕士，量化研究经历覆盖数据清洗、因子/策略回测、财报与公告采集；能讲清回测口径和失效边界。",
        "现职从 0 搭建部门 AI 投研流程：日报工作流、知识库对接、盘中信号推送、Linux 定时任务。",
        "用 Cursor / Claude / MCP 把投研步骤做成可重复 Skill（采集→校验→失败换源），不是只会对话。",
        "Xpert 金融 Skill 评测、TalentsAI Agentic Coding × 金融已过审。",
    ]:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.left_indent = Cm(0.6)
        set_run_font(p.add_run(line), size=10)

    add_heading_bar(doc, "工作经历")
    add_job(
        doc,
        "湖北正知资产管理有限公司（私募，中基协 P1007881） | AI投研 / 量化研究",
        "2026.06 – 至今",
        [
            "部门 AI 投研从 0 搭建：对比 WorkBuddy 与 Coze、配置工作空间；部署投研智能体并接通微信/飞书。",
            "Linux 部署脚本与定时任务；日报工作流多人共用、企微推送与自动归档；乐享知识库对接 Coze。",
            "盘中放量扫描与股票池多时段推送；强势股策略加「收盘价/前20日最低价<130%」过滤，对照回测后最大单笔亏损约 -27% 降至约 -17%。",
        ],
    )
    add_job(
        doc,
        "量化研究工作室 | 独立研究员",
        "2024.03 – 2026.04",
        [
            "结合宏观与行业框架研究热点赛道轮动，为策略定调提供参考；搭建财务选股模型，按季度更新核心股票池。",
            "建设全市场财报采集与业绩预告/快报爬虫；结合短线指标完善买卖点逻辑；日常用 Claude、Cursor、MCP 提效。",
        ],
    )
    add_job(
        doc,
        "江铜产融控股有限公司 | 量化研究员",
        "2023.10 – 2024.02",
        [
            "对北上深量化私募做调研问卷、实地尽调并撰写尽调报告；搭建二级市场研究框架，监控流动性等指标。",
        ],
    )
    add_job(
        doc,
        "量盈私募基金管理有限公司 | 量化研究员",
        "2022.01 – 2023.09",
        [
            "用聚宽、掘金做期货/股票因子挖掘、策略开发与回测优化；按基金经理要求用 Python 复现券商金工研报。",
        ],
    )
    add_job(
        doc,
        "深圳引力波量化科技有限公司 | 量化研究员",
        "2020.08 – 2021.12",
        [
            "交易数据获取，用 Pandas、Numpy、Talib 清洗与分析；参与套利、做市等策略开发与日常维护。",
        ],
    )
    add_job(
        doc,
        "中国平安财产保险股份有限公司 | 再保险业务管理岗（管培生）",
        "2019.07 – 2020.07",
        [
            "Excel 做再保人资信月度更新、季度坏账计提与风险指标监控；汇总关联/非关联交易分类数据。",
        ],
    )

    add_heading_bar(doc, "项目经历")
    add_job(
        doc,
        "AI 投研工作流与公告采集 Agent",
        "2026.06 – 至今",
        [
            "日报工作流技能化、知识库对接、盘中信号推送；Linux 部署与定时任务。",
            "巨潮公告采集：固定步骤→校验→失败换源。LangGraph demo：github.com/invy95/myCode（langgraph-announcement-agent 分支）。",
        ],
    )
    add_job(
        doc,
        "股票短线多头（二级市场研究）",
        "2024.03 – 至今",
        [
            "赛道研究、财报采集、业绩预告爬虫。个人账户：2025 年资金加权年化约 60%，2026 年约 40%（与 11 号简历一致）。",
        ],
    )
    add_job(
        doc,
        "股票短线多头（掘金）",
        "2022.07 – 至今",
        [
            "分钟级板块因子 + 日线量价择时选股；实盘版接入掘金日更数据。样本内：年化 17.14%，最大回撤 2%，夏普 3.4。",
        ],
    )
    add_job(
        doc,
        "期货 CTA 模式识别",
        "2022.01 – 2022.03",
        [
            "Python + tbquant，基于拓扑空间结构距离；股指期货 2012–2018 分年回测。",
        ],
    )

    add_heading_bar(doc, "教育背景")
    add_job(
        doc,
        "华中科技大学 | 金融 | 硕士（全日制）",
        "2017.09 – 2019.06",
        [
            "主修量化投资、财务报表分析、金融建模与计算、数理统计。考研 405/500（数学 130，英语 79，专业课 134）。",
            "论文：《医药行业上市公司价值评估研究——以云南白药为例》。",
        ],
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run("武汉科技大学 | 机械工程 | 本科（全日制）")
    set_run_font(r1, size=10.5, bold=True)
    r2 = p.add_run("    2013.09 – 2017.06")
    set_run_font(r2, size=10, color=RGBColor(0x55, 0x55, 0x55))

    add_heading_bar(doc, "技能证书")
    for line in [
        "数据：Python（Pandas / Numpy / Talib）、Stata、MATLAB、Wind、Excel。",
        "工程：Linux、脚本部署、cron、日志；聚宽、掘金。",
        "AI：Claude、Cursor、MCP、Coze；能把投研步骤做成 Skill（校验+换源）。",
        "近期：Xpert 金融 Skill 评测已过审；TalentsAI Agentic Coding × 金融已过审。",
        "证书：证券 / 基金 / 银行从业、计算机二级；英语六级 558。",
    ]:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.left_indent = Cm(0.6)
        set_run_font(p.add_run(line), size=10)

    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
