# 功能：按本人上传的中文简历《殷维简历_15671678098.docx》生成英文版。
# 更新：2026-09-13：全文英译；个人作品集四条网址与中文稿完全一致，不做改写或纠错。

"""输出 YinWei_Resume_EN.docx。结构与中文稿一致，目标尽量两页。"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path(__file__).resolve().parent / "YinWei_Resume_EN.docx"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GRAY = RGBColor(0x55, 0x55, 0x55)
LINK_BLUE = RGBColor(0x05, 0x63, 0xC1)
BODY = 9
LINE = 1.0
FONT = "Calibri"

# 与中文稿逐字相同，禁止改动（含 minduful 拼写与查询参数）。
PORTFOLIO = [
    (
        "Financial investment-research Skills: ",
        "https://github.com/invy95/myCode/raw/cursor/talentsai-agentic-6afc/resume/talentsai/TalentsAI_AgenticCoding%E9%87%91%E8%9E%8D_%E6%9C%BA%E4%BC%9A%E8%83%9C%E4%BB%BB%E4%B8%8E%E8%83%BD%E5%8A%9B%E8%AF%81%E6%98%8E.docx",
    ),
    (
        "Breathing Bubble meditation desktop widget: ",
        "https://github.com/invy95/minduful-breathing/releases/download/v1.0.1/BreathingBall-Open.zip",
    ),
    (
        "AI explainer video: ",
        "https://m.bilibili.com/video/BV1E9Y46oEGX?share_source=weixin_web&vd_source=9911c2ba7a4191da6df1074d816fc961",
    ),
    (
        "Douyin video: ",
        "https://v.douyin.com/Hszw0uTsQuY/",
    ),
]


def set_run_font(run, name=FONT, size=BODY, bold=False, color=None, underline=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if color is not None:
        run.font.color.rgb = color
    if underline is not None:
        run.underline = underline


def _set_bool(p_pr, tag, val="0"):
    el = p_pr.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        p_pr.append(el)
    el.set(qn("w:val"), val)


def tight(p, before=0, after=0, line=LINE):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p_pr = p._p.get_or_add_pPr()
    _set_bool(p_pr, "w:widowControl", "0")
    _set_bool(p_pr, "w:autoSpaceDE", "0")
    _set_bool(p_pr, "w:autoSpaceDN", "0")
    _set_bool(p_pr, "w:adjustRightInd", "0")
    _set_bool(p_pr, "w:snapToGrid", "0")


def add_center(doc, text, size, bold=False, after=0, color=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tight(p, 0, after)
    set_run_font(p.add_run(text), size=size, bold=bold, color=color)
    return p


def add_section(doc, text):
    p = doc.add_paragraph()
    tight(p, 3, 1)
    set_run_font(p.add_run(text), size=10.5, bold=True, color=NAVY)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F3A5F")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)
    return p


def add_line(doc, text, size=BODY, bold=False, before=0, after=0, color=None):
    p = doc.add_paragraph()
    tight(p, before, after)
    set_run_font(p.add_run(text), size=size, bold=bold, color=color)
    return p


def add_hyperlink(paragraph, url, size=BODY):
    """显示文字 = 传入网址本身，关系目标也用同一字符串，避免改链。"""
    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT)
    r_fonts.set(qn("w:hAnsi"), FONT)
    r_fonts.set(qn("w:eastAsia"), "微软雅黑")
    r_pr.append(r_fonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    r_pr.append(sz)
    sz_cs = OxmlElement("w:szCs")
    sz_cs.set(qn("w:val"), str(int(size * 2)))
    r_pr.append(sz_cs)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(u)
    new_run.append(r_pr)
    text = OxmlElement("w:t")
    text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    text.text = url
    new_run.append(text)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_link_line(doc, label, url, before=0):
    p = doc.add_paragraph()
    tight(p, before, 0)
    set_run_font(p.add_run(label), size=8.5)
    add_hyperlink(p, url, size=8.5)
    return p


def add_entry_head(doc, title, dates, subtitle=None):
    p = doc.add_paragraph()
    tight(p, 3, 0)
    set_run_font(p.add_run(title), size=BODY, bold=True)
    if subtitle:
        set_run_font(p.add_run("  |  " + subtitle), size=BODY, color=GRAY)
    set_run_font(p.add_run("    " + dates), size=9, color=GRAY)
    return p


def add_job(doc, company, dates, title, bullets):
    add_entry_head(doc, company, dates, title)
    for i, b in enumerate(bullets, start=1):
        add_line(doc, f"{i}. {b}")


def add_project(doc, name, dates, desc, duties, result=None):
    add_entry_head(doc, "● " + name, dates)
    add_line(doc, "Description: " + desc)
    for i, b in enumerate(duties, start=1):
        add_line(doc, f"{i}. {b}")
    if result:
        add_line(doc, "Results: " + result)


def main():
    doc = Document()
    for sec in doc.sections:
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.top_margin = Cm(1.05)
        sec.bottom_margin = Cm(1.05)
        sec.left_margin = Cm(1.35)
        sec.right_margin = Cm(1.35)
        sect_pr = sec._sectPr
        doc_grid = sect_pr.find(qn("w:docGrid"))
        if doc_grid is None:
            doc_grid = OxmlElement("w:docGrid")
            sect_pr.append(doc_grid)
        doc_grid.set(qn("w:type"), "lines")
        doc_grid.set(qn("w:linePitch"), "240")

    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(BODY)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    s_pr = style.element.get_or_add_pPr()
    _set_bool(s_pr, "w:widowControl", "0")
    _set_bool(s_pr, "w:autoSpaceDE", "0")
    _set_bool(s_pr, "w:autoSpaceDN", "0")

    add_center(doc, "Yin Wei", 16, bold=True, after=1, color=NAVY)
    add_center(
        doc,
        "Female  |  Age 31  |  CPC Member  |  Wuhan, Hubei  |  Master's in Finance  |  15671678098  |  yw43@foxmail.com",
        9.5,
        color=GRAY,
    )

    add_section(doc, "Work Experience")
    add_job(
        doc,
        "Hubei Zhengzhi Asset Management Co., Ltd. (private fund; AMAC P1007881)",
        "Jun 2026 – Present",
        "AI Investment Research Manager",
        [
            "Built the investment-research workspace around workflows and a knowledge base to support the team's AI rollout. Compared Coze and other AI tools across multiple devices and clients, selected the stack, and launched Investment-Research Workflow 1.0 and Quantitative Trading System 1.0 in the research setting, including automated single-stock research notes.",
            "Sent market-data updates to research colleagues: scanned for unusual intraday volume and pushed watchlist highlights at the morning, midday, and afternoon close sessions to support trading decisions.",
            'Produced an "AI retreat report" on tech-versus-growth style rotations, connected it to the Wind API, and sent an automatically refreshed daily brief.',
        ],
    )
    add_job(
        doc,
        "Quantitative Research Studio",
        "Mar 2024 – Jun 2026",
        "Secondary-Market Researcher",
        [
            "Combined macroeconomic and industry frameworks to study thematic rotation and supplied the logic used to set strategy tone.",
            "Built a standardized fundamental stock-selection AI model, constructed a quarterly core universe, and refined entry and exit references with short-term technical indicators.",
        ],
    )
    add_job(
        doc,
        "Jiangxi Copper Financial Holdings Co., Ltd.",
        "Oct 2023 – Feb 2024",
        "Investment Management",
        [
            "Designed questionnaires, conducted on-site due diligence on quantitative managers in Beijing, Shanghai, Guangzhou, and Shenzhen, and wrote manager due-diligence reports.",
        ],
    )
    add_job(
        doc,
        "Liangying Private Fund Management Co., Ltd.",
        "Jan 2022 – Sep 2023",
        "Quantitative Researcher",
        [
            "Mined futures and equity factors, developed and backtested strategies, and reviewed the results for further optimization.",
            "Replicated sell-side financial-engineering strategy reports in Python at the portfolio manager's request.",
        ],
    )
    add_job(
        doc,
        "Shenzhen Gravity Wave Quantitative Technology Co., Ltd.",
        "Aug 2020 – Dec 2021",
        "Market-Making Product Manager",
        [
            "Worked directly with the exchange. Translated order-book display requirements into executable parameters: price placement from bid one through bid five, size at each level, and the interval in seconds for sequential quoting.",
            "Implemented and launched the quoting program in Python on the exchange API; built physical and strategy-level risk controls; monitored the book and, in extreme markets, canceled orders and flattened positions by trading in the opposite direction.",
            "Retrieved and cleaned trading data (Pandas, NumPy, TA-Lib) and tracked arbitrage P&L on the market-making project.",
        ],
    )
    add_job(
        doc,
        "Ping An Property & Casualty Insurance Company of China, Ltd.",
        "Jul 2019 – Jul 2020",
        "Reinsurance Administration (Management Trainee)",
        [
            "Used Excel to update and review reinsurer credit files monthly, book quarterly bad-debt provisions, monitor risk indicators, and consolidate related-party and non-related-party transaction data.",
        ],
    )

    add_section(doc, "Education")
    add_line(
        doc,
        "Sep 2017 – Jun 2019    Huazhong University of Science and Technology    Finance    Master's (full-time)",
        before=1,
    )
    add_line(
        doc,
        "Sep 2013 – Jun 2017    Wuhan University of Science and Technology    Mechanical Engineering    Bachelor's (full-time)",
    )

    add_section(doc, "Projects")
    add_project(
        doc,
        "Investment-Research AI: Turning Requests into Decision Support",
        "Jun 2026 – Present",
        "Used AI to turn research-note generation, announcements, and market scans into runnable workflows, and turned colleagues' requests into reusable processes and systems.",
        [
            "Built the team's research platform—daily workflows, AI research workflows, and a knowledge base—and connected the workspace end to end. Helped the team keep an AI-supported audit trail, archive work, collaborate, and conduct research.",
            "Deployed AI-generated research notes and market reviews to a research agent and pushed them to WeChat and Feishu each day.",
            "Parsed market data (intraday volume-spike scans) and pushed the watchlist by session.",
        ],
    )
    add_project(
        doc,
        "Equity Investment Research",
        "Mar 2024 – Present",
        "Used Cursor and other AI agents to build a multi-factor stock-analysis framework covering fundamentals, sentiment, and technicals: parsing market data, organizing the logic, and attributing earnings to support strategy research.",
        [
            "Built a filings-data pipeline that pulls A-share fundamentals through Tushare and computes related factor values.",
            "Wrote crawlers for earnings pre-announcements and flash reports, and designed a real-time hot-concept dashboard.",
        ],
        "The work deepened understanding of fundamentals, short-term technical patterns, and sentiment cycles, and turned research ideas into checkable stock-selection rules.",
    )
    add_project(
        doc,
        "Short-Term Long-Only Equity Strategy",
        "Jul 2022 – Feb 2023",
        "Developed a short-term long-only equity strategy on the GoldMiner platform. Combined minute-level sector factors with daily price-volume factors for timing and selection, then ran the strategy in live trading. Mined six factors; a three-factor combination performed best.",
        [
            "Combined price-volume relationships into timing and selection factors, designed entry and exit rules, and backtested parameter and factor combinations.",
            "Wrote a live-trading version and connected it to GoldMiner's daily data updates.",
        ],
        "In-sample backtest: 2% maximum drawdown, 17.14% annualized return, and a Sharpe ratio of 3.4.",
    )
    add_project(
        doc,
        "Financial-Engineering Report Replication",
        "Nov 2022 – Jan 2023",
        "Replicated the financial-engineering report “3M Sector Rotation Strategy” in Python using the Wind database.",
        [
            "Studied the report, collected the relevant micro-level data from Wind, and implemented the factor algorithms.",
        ],
        "Micro-factor data collection and cross-sectional factor outputs were completed.",
    )
    add_project(
        doc,
        "Futures CTA Strategy Research",
        "Jan 2022 – Mar 2022",
        "Developed and optimized a futures pattern-recognition strategy in Python with TBQuant.",
        [
            "Connected TBQuant to Python for data processing and handoff, and built a pattern-recognition algorithm based on topological spatial-structure distance.",
            "Applied the strategy to index futures and refined the backtest, the strategy, and its parameters.",
        ],
        "An index-futures backtest from 2012 through 2018 was positive in each of the seven years, with an annualized return of 80%.",
    )
    add_project(
        doc,
        "Market-Making Project: Requirements → Quoting Design → Launch",
        "Dec 2020 – Dec 2021",
        "Provided liquidity to an exchange, covering requirements, design, and implementation end to end.",
        [
            "Aligned with the exchange on liquidity targets and order-book shape, and designed quoting-sequence parameters so that five-level prices, sizes, and stagger timing stayed close to a live market.",
            "Implemented and launched the program in Python; paired physical and programmatic risk controls; monitored extreme-market risk and manually reversed to flat when needed.",
        ],
        "Increased trading activity on the exchange and completed the loop from requirements to parameter design, implementation, launch, and monitoring.",
    )

    add_section(doc, "Honors & Skills")
    add_line(
        doc,
        "Securities, fund, and banking qualification certificates; National Computer Rank Examination Level 2; CET-4 and CET-6 (CET-6 score 558). Master's thesis: “Valuation of Listed Companies in the Pharmaceutical Industry: A Case Study of Yunnan Baiyao” (Gordon dividend-discount model, with a quantitative strategy written in Python for validation).",
    )
    add_line(
        doc,
        "I am proficient in Python, MATLAB, Wind, Stata, Excel, Access, and PowerPoint, and I can combine Wind and Excel for financial data analysis.",
    )
    add_line(
        doc,
        "Daily work includes using Cursor and Claude to break down, implement, and review investment-research tasks; connecting tools into workflows through MCP; building research agents in Coze and connecting them to WeChat and Feishu; encoding multi-step tasks as runnable agent flows (LangGraph); and writing skills for the Coze store. I am familiar with the use cases of different AI tools and models, and I can write code independently or use AI to write code for lightweight checks.",
    )
    add_line(
        doc,
        "Recent remote work includes the Xpert financial-expert AI annotation project and quantitative-studio collaboration delivered remotely.",
    )

    add_section(doc, "Portfolio")
    for i, (label, url) in enumerate(PORTFOLIO):
        add_link_line(doc, label, url, before=1 if i == 0 else 0)

    body = doc.element.body
    for child in list(body):
        if child.tag == qn("w:p") and not "".join(child.itertext()).strip():
            p_pr = child.find(qn("w:pPr"))
            if p_pr is not None and p_pr.find(qn("w:sectPr")) is not None:
                continue
            if child.getnext() is None:
                continue
            body.remove(child)

    doc.save(OUT)
    print(f"Generated: {OUT}")


if __name__ == "__main__":
    main()
