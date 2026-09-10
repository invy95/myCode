# 功能：把原版 PDF《殷维_简历1(1).pdf》（WPS 双栏模板）做成可编辑 Word，版式对齐原件。
# 更新：2026-09-10：按 PDF 原结构生成两页 A4 Word，不改正文事实。

"""输出 resume/殷维_简历.docx。"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
PDF = Path("/workspace/殷维_简历1(1).pdf")
PHOTO = Path("/tmp/orig-pdf/img-p1-0.png")
OUT = ROOT / "殷维_简历.docx"

NAVY = RGBColor(0x41, 0x70, 0x9C)
DARK = RGBColor(0x2C, 0x3D, 0x50)
GRAY = RGBColor(0x4A, 0x4A, 0x4A)
BLACK = RGBColor(0x22, 0x22, 0x22)
BANNER_HEX = "2C3D50"
BAR_HEX = "2C3D50"
BAR_EMPTY_HEX = "E8E8E8"


def set_run_font(run, name="微软雅黑", size=10, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color is not None:
        run.font.color.rgb = color


def tight(p, before=0, after=0, line=1.05):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_margins(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for key, val_cm in kwargs.items():
        node = OxmlElement(f"w:{key}")
        node.set(qn("w:w"), str(int(val_cm * 567)))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def no_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    tblPr.append(borders)


def set_table_width(table, width_cm):
    table.autofit = False
    table.allow_autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(int(width_cm * 567)))
    tblW.set(qn("w:type"), "dxa")


def add_para(cell, text, size=10, bold=False, color=BLACK, before=0, after=1, line=1.08):
    p = cell.add_paragraph()
    tight(p, before, after, line)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def banner(cell, text):
    table = cell.add_table(rows=1, cols=1)
    no_borders(table)
    c = table.cell(0, 0)
    shade_cell(c, BANNER_HEX)
    set_cell_margins(c, top=0.08, bottom=0.08, left=0.18, right=0.1)
    p = c.paragraphs[0]
    tight(p, 0, 0, 1.0)
    run = p.add_run(text)
    set_run_font(run, size=11, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))


def kv(cell, label, value):
    p = cell.add_paragraph()
    tight(p, 1, 1, 1.08)
    r1 = p.add_run(f"{label}")
    set_run_font(r1, size=9.5, color=GRAY)
    r2 = p.add_run(value)
    set_run_font(r2, size=9.5, color=BLACK)


def set_row_height(row, pt):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    trHeight = OxmlElement("w:trHeight")
    trHeight.set(qn("w:val"), str(int(pt * 20)))
    trHeight.set(qn("w:hRule"), "exact")
    trPr.append(trHeight)


def skill_line(cell, name, pct):
    p = cell.add_paragraph()
    tight(p, 1, 1, 1.0)
    filled = max(1, min(10, int(round(10 * pct))))
    bar = "■" * filled + "□" * (10 - filled)
    set_run_font(p.add_run(f"{name}  "), size=9, color=GRAY)
    set_run_font(p.add_run(bar), size=8, color=DARK)


def section_title(cell, text, first=False):
    p = cell.add_paragraph() if not first else cell.paragraphs[0]
    if first:
        p = cell.paragraphs[0]
    else:
        p = cell.add_paragraph()
    tight(p, 8 if not first else 0, 4, 1.0)
    run = p.add_run(text)
    set_run_font(run, size=13, bold=True, color=NAVY)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "41709C")
    pBdr.append(bottom)
    pPr.append(pBdr)


def job_block(cell, dates, company, title, bullets):
    p = cell.add_paragraph()
    tight(p, 8, 0, 1.05)
    r1 = p.add_run(dates)
    set_run_font(r1, size=9.5, color=GRAY)
    r2 = p.add_run("    " + company)
    set_run_font(r2, size=10.5, bold=True, color=DARK)
    p2 = cell.add_paragraph()
    tight(p2, 0, 2, 1.05)
    r3 = p2.add_run(title)
    set_run_font(r3, size=9.5, bold=True, color=NAVY)
    for b in bullets:
        p = cell.add_paragraph()
        tight(p, 0, 1, 1.08)
        run = p.add_run(b)
        set_run_font(run, size=9.5, color=BLACK)


def project_block(cell, name, dates, desc, duties, result, first=False):
    p = cell.add_paragraph() if not first else cell.paragraphs[0]
    if first:
        # title already used
        p = cell.add_paragraph()
    tight(p, 8, 1, 1.05)
    r0 = p.add_run("●  " + name)
    set_run_font(r0, size=11, bold=True, color=DARK)
    tab = p.add_run("    " + dates)
    set_run_font(tab, size=9.5, color=GRAY)
    add_para(cell, "项目描述：" + desc, size=9.5, after=1)
    add_para(cell, "个人职责：", size=9.5, after=0)
    for b in duties:
        add_para(cell, b, size=9.5, after=0)
    add_para(cell, "项目业绩：" + result, size=9.5, after=4)


def prevent_row_split(row):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trPr.append(cant)


def main():
    if not PHOTO.exists():
        import pymupdf

        doc = pymupdf.open(PDF)
        xref = doc[0].get_images(full=True)[0][0]
        pix = pymupdf.Pixmap(doc, xref)
        if pix.n >= 5:
            pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
        PHOTO.parent.mkdir(parents=True, exist_ok=True)
        pix.save(str(PHOTO))

    doc = Document()
    for sec in doc.sections:
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.top_margin = Cm(1.2)
        sec.bottom_margin = Cm(1.2)
        sec.left_margin = Cm(1.3)
        sec.right_margin = Cm(1.3)

    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(10)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    usable = 21.0 - 1.3 - 1.3
    left_w, right_w = 6.2, usable - 6.2

    table = doc.add_table(rows=1, cols=2)
    no_borders(table)
    set_table_width(table, usable)
    table.columns[0].width = Cm(left_w)
    table.columns[1].width = Cm(right_w)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    left, right = table.cell(0, 0), table.cell(0, 1)
    set_cell_margins(left, top=0.05, bottom=0.05, left=0.05, right=0.25)
    set_cell_margins(right, top=0.05, bottom=0.05, left=0.25, right=0.05)
    prevent_row_split(table.rows[0])

    p = left.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tight(p, 0, 4)
    p.add_run().add_picture(str(PHOTO), width=Cm(3.3))

    p = left.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tight(p, 2, 6)
    set_run_font(p.add_run("殷维"), size=16, bold=True, color=NAVY)

    banner(left, "个人资料")
    kv(left, "性    别：", "女")
    kv(left, "年    龄：", "31 岁")
    kv(left, "籍    贯：", "湖北武汉")
    kv(left, "政治面貌：", "中共党员")
    kv(left, "专    业：", "金融")
    kv(left, "学    历：", "硕士")
    kv(left, "身    高：", "165cm")

    p = left.add_paragraph()
    tight(p, 8, 4)
    banner(left, "联系方式")
    kv(left, "手机：", "15671678098")
    kv(left, "邮箱：", "yw43@foxmail.com")
    kv(left, "地址：", "武汉市青山区")

    p = left.add_paragraph()
    tight(p, 8, 4)
    banner(left, "掌握技能")
    p = left.add_paragraph()
    tight(p, 4, 1)
    set_run_font(p.add_run("技能"), size=9.5, bold=True, color=DARK)
    skill_line(left, "Python:", 0.86)
    skill_line(left, "Stata：", 0.80)
    skill_line(left, "Excel：", 0.68)
    p = left.add_paragraph()
    tight(p, 6, 1)
    set_run_font(p.add_run("语言"), size=9.5, bold=True, color=DARK)
    skill_line(left, "普通话：", 0.90)
    skill_line(left, "英语：", 0.52)

    section_title(right, "教育背景", first=True)
    add_para(right, "2017.09-2019.06    华中科技大学    金融    硕士（全日制）", size=10, before=4, after=1)
    add_para(right, "2013.09-2017.06    武汉科技大学    机械工程    本科（全日制）", size=10, after=2)

    section_title(right, "主修课程")
    add_para(right, "考研成绩：总分（405/500）数学（130/150）英语（79/100）专业课（134/150）", size=9.5, before=4, after=1)
    add_para(right, "主修课程：量化投资、财务报表分析、金融建模与计算、数理统计等", size=9.5, after=2)

    section_title(right, "工作经历")
    job_block(
        right,
        "2024.03-2026.06",
        "量化研究工作室",
        "独立研究员",
        [
            "1、结合宏观与行业研究框架，研究热点赛道轮动，为策略定调提供参考。",
            "2、搭建标准化财务选股模型，构建季度更新核心股票池，精准筛选优质个股标的。",
            "3、结合短线指标，完善个股买卖点参考逻辑。",
        ],
    )
    job_block(
        right,
        "2023.10-2024.02",
        "江铜产融控股有限公司",
        "量化研究员",
        [
            "1、对北上深量化私募进行调研前问卷设计、实地尽调，并根据调研情况撰写尽调报告。",
            "2、搭建二级市场研究框架，监控流动性等。",
        ],
    )
    job_block(
        right,
        "2022.01-2023.09",
        "量盈私募基金管理有限公司",
        "量化研究员",
        [
            "1、使用聚宽、掘金等常见量化平台进行期货、股票因子挖掘、策略开发及回测优化。",
            "2、根据基金经理要求，使用 python 对券商金工等策略研报进行复现。",
        ],
    )
    job_block(
        right,
        "2020.08-2021.12",
        "深圳引力波量化科技有限公司",
        "量化研究员",
        [
            "1、交易数据获取，熟练使用 Pandas、Numpy、Talib 等 package 对数据进行清洗、分析。",
            "2、量化策略开发等工作，包括套利、做市等方向，日常优化、维护交易策略。",
        ],
    )
    job_block(
        right,
        "2019.07-2020.07",
        "中国平安财产保险股份有限公司",
        "再保险业务管理岗（管培生）",
        [
            "1、熟练使用 excel 对再保人资信进行月度更新、季度坏账计提、监控风险指标。",
            "2、根据再保险资信管理制度，通过市场及同业信息，审核交易对手资信情况。",
            "3、统计汇总不同业务场景分类下关联交易、非关联交易情况数据。",
        ],
    )

    doc.add_page_break()

    p = doc.add_paragraph()
    tight(p, 0, 4, 1.0)
    run = p.add_run("项目经历")
    set_run_font(run, size=13, bold=True, color=NAVY)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "41709C")
    pBdr.append(bottom)
    pPr.append(pBdr)

    def add_project(name, dates, desc, duties, result):
        p = doc.add_paragraph()
        tight(p, 10, 2, 1.05)
        r0 = p.add_run("●  " + name)
        set_run_font(r0, size=11, bold=True, color=DARK)
        r1 = p.add_run("                    " + dates)
        set_run_font(r1, size=9.5, color=GRAY)
        for line, kw in [
            ("项目描述：" + desc, False),
            ("个人职责：", False),
            *[(d, False) for d in duties],
            ("项目业绩：" + result, False),
        ]:
            p = doc.add_paragraph()
            tight(p, 0, 1, 1.08)
            set_run_font(p.add_run(line), size=10, color=BLACK)

    add_project(
        "股票短线多头策略",
        "2024.03-至今",
        "二级市场研究",
        [
            "1、对宏观赛道进行行业研究，出具研究报告。",
            "2、搭建财报数据收集系统收集全市场个股基本面财报数据。",
            "3、编写业绩预报、快报爬虫，及时爬取个股相关业绩信息。",
        ],
        "取得较好的投资收益，并积累了较好的盘感及重度使用 AI 工具辅助任务的实践经验。2025 年资金加权收益率年化 60%，2026 年资金加权收益率年化 40%。",
    )
    add_project(
        "股票短线多头策略",
        "2022.07-至今",
        "使用掘金平台开发股票短线多头策略，通过分钟级别板块因子组合日线级别量价因子进行择时选股获得超额收益。",
        [
            "1、通过分析量价关系，组合合适的择时、选股因子，结合入场、出场方式的选择，研发股票策略。",
            "2、对策略结果进行回测，对参数、因子组合进行优化。",
            "3、编写实盘策略版本，接入掘金每日更新数据。",
        ],
        "回测样本内最大回撤 2%，年化收益 17.14%，夏普比率 3.4。",
    )
    add_project(
        "期货 cta 策略研究",
        "2022.01-2022.03",
        "使用 python 结合 tbquant 针对“基于期货的模式识别”进行策略开发及优化工作。",
        [
            "1、基于拓扑空间结构距离开发模式识别策略算法开发。",
            "2、模式识别策略应用于股指期货市场的回测情况进行优化。",
        ],
        "基于股指期货回测 2012-2018 年连续 7 年每年回测收益率为正，年化收益率 80%。",
    )

    p = doc.add_paragraph()
    tight(p, 12, 4, 1.0)
    run = p.add_run("荣誉技能")
    set_run_font(run, size=13, bold=True, color=NAVY)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "41709C")
    pBdr.append(bottom)
    pPr.append(pBdr)

    for line in [
        "证券从业资格证书、基金从业资格证书、银行从业资格证书",
        "计算机二级证书，熟练掌握 Stata、Excel、Access、PPT 等软件",
        "通过英语四、六级（558），快速阅读英文文献，交流能力良好",
        "熟练掌握 python 语言、matlab、wind 等软件，能使用 wind 结合 excel 进行金融数据分析",
        "基于股利贴现模型和量化分析，完成硕士毕业论文《医药行业上市公司价值评估研究--以云南白药为例》",
        "熟练使用前沿 AI 工具和模型，如 claude、cursor 等高效完成任务，了解 MCP 基本用法，能搭建 AI agent team 协同提效。",
    ]:
        p = doc.add_paragraph()
        tight(p, 1, 2, 1.12)
        set_run_font(p.add_run(line), size=10.5, color=BLACK)

    doc.save(OUT)
    print(f"已生成：{OUT}")


if __name__ == "__main__":
    main()
